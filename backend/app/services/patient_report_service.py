from __future__ import annotations

import html
import io
from datetime import UTC, datetime
from decimal import Decimal, ROUND_HALF_UP

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models import (
    Appointment,
    CNSSCase,
    DoctorProfile,
    Fee,
    MedicalDocument,
    PatientCabinetMembership,
    PatientProfile,
    Payment,
    Prescription,
    Tenant,
    Visit,
)


INVALID_PAYMENT_STATUSES = {
    "ANNULE", "REMBOURSE", "ECHEC", "FAILED", "REVERSED", "CANCELLED",
}
PAYMENT_LABELS = {
    "cash": "Espèces", "card": "Carte bancaire", "transfer": "Virement", "cheque": "Chèque",
}


def _decimal(value: object) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _money(value: object) -> str:
    return f"{_decimal(value):,.2f} MAD".replace(",", " ").replace(".", ",")


def _valid_payment(item: Payment) -> bool:
    return item.status not in INVALID_PAYMENT_STATUSES and (
        item.method != "cheque" or item.cheque_status == "ENCAISSE"
    )


def patient_report_data(db: Session, tenant_id: str, patient_id: str) -> dict:
    tenant = db.scalar(select(Tenant).where(Tenant.id == tenant_id))
    patient = db.scalar(
        select(PatientProfile)
        .options(joinedload(PatientProfile.user))
        .where(PatientProfile.id == patient_id)
    )
    membership = db.scalar(select(PatientCabinetMembership.id).where(
        PatientCabinetMembership.patient_id == patient_id,
        PatientCabinetMembership.tenant_id == tenant_id,
        PatientCabinetMembership.status == "active",
    )) if patient else None
    if not tenant or not patient or (patient.tenant_id != tenant_id and not membership):
        raise HTTPException(404, "Patient introuvable dans ce cabinet.")

    appointments = list(db.scalars(select(Appointment).where(
        Appointment.tenant_id == tenant_id, Appointment.patient_id == patient_id,
    ).order_by(Appointment.date.desc(), Appointment.start_time.desc())).all())
    visits = list(db.scalars(select(Visit).where(
        Visit.tenant_id == tenant_id, Visit.patient_id == patient_id,
    ).order_by(Visit.visit_date.desc())).all())
    fees = list(db.scalars(select(Fee).where(
        Fee.tenant_id == tenant_id, Fee.patient_id == patient_id,
    ).order_by(Fee.created_at.desc())).all())
    payments = list(db.scalars(select(Payment).where(
        Payment.tenant_id == tenant_id, Payment.patient_id == patient_id,
    ).order_by(Payment.paid_at.desc())).all())
    prescriptions = list(db.scalars(select(Prescription).where(
        Prescription.tenant_id == tenant_id, Prescription.patient_id == patient_id,
    ).order_by(Prescription.created_at.desc())).all())
    documents = list(db.scalars(select(MedicalDocument).where(
        MedicalDocument.tenant_id == tenant_id,
        MedicalDocument.patient_id == patient_id,
        MedicalDocument.status != "ARCHIVE",
    ).order_by(MedicalDocument.created_at.desc())).all())
    cnss_records = list(db.scalars(select(CNSSCase).where(
        CNSSCase.tenant_id == tenant_id, CNSSCase.patient_id == patient_id,
    ).order_by(CNSSCase.created_at.desc())).all())

    doctor_ids = {item.doctor_id for item in [*appointments, *visits, *prescriptions] if item.doctor_id}
    doctors = {
        item.id: f"Dr {item.user.first_name} {item.user.last_name}".strip()
        for item in db.scalars(
            select(DoctorProfile).options(joinedload(DoctorProfile.user)).where(DoctorProfile.id.in_(doctor_ids))
        ).all()
    } if doctor_ids else {}
    valid_payments = [item for item in payments if _valid_payment(item)]
    payments_by_fee: dict[str, Decimal] = {}
    for payment in valid_payments:
        payments_by_fee[payment.fee_id] = payments_by_fee.get(payment.fee_id, Decimal("0")) + _decimal(payment.amount)
    fee_rows = []
    for fee in fees:
        total = _decimal(fee.amount)
        paid = _decimal(payments_by_fee.get(fee.id, 0))
        fee_rows.append({"item": fee, "total": total, "paid": paid, "remaining": max(Decimal("0"), total - paid)})
    active_fee_ids = {row["item"].id for row in fee_rows if row["item"].status != "cancelled"}
    for row in fee_rows:
        if row["item"].id not in active_fee_ids:
            row["remaining"] = Decimal("0")
    total_fees = sum((row["total"] for row in fee_rows if row["item"].id in active_fee_ids), Decimal("0"))
    total_paid = sum((_decimal(item.amount) for item in valid_payments if item.fee_id in active_fee_ids), Decimal("0"))
    return {
        "tenant": tenant,
        "patient": patient,
        "patient_name": f"{patient.user.first_name} {patient.user.last_name}".strip(),
        "appointments": appointments,
        "visits": visits,
        "fees": fee_rows,
        "payments": payments,
        "prescriptions": prescriptions,
        "documents": documents,
        "cnss_records": cnss_records,
        "doctors": doctors,
        "total_fees": total_fees,
        "total_paid": total_paid,
        "rap": max(Decimal("0"), total_fees - total_paid),
        "generated_at": datetime.now(UTC),
    }


def render_patient_pdf(data: dict) -> bytes:
    output = io.BytesIO()
    tenant, patient = data["tenant"], data["patient"]
    brand = colors.HexColor(getattr(tenant, "primary_color", None) or "#0f766e")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("PatientTitle", parent=styles["Title"], textColor=brand, fontSize=19, leading=23)
    section = ParagraphStyle("PatientSection", parent=styles["Heading2"], textColor=colors.HexColor("#0b3441"), fontSize=12, leading=15, spaceBefore=7, spaceAfter=5)
    cell = ParagraphStyle("PatientCell", parent=styles["Normal"], fontSize=7.5, leading=9.5)
    right = ParagraphStyle("PatientMoney", parent=cell, alignment=TA_RIGHT)
    center = ParagraphStyle("PatientCenter", parent=cell, alignment=TA_CENTER)

    def p(value: object, style=cell) -> Paragraph:
        return Paragraph(html.escape(str(value if value not in (None, "") else "—")), style)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#dbe8e5"))
        canvas.line(14 * mm, 11 * mm, 283 * mm, 11 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(14 * mm, 6.5 * mm, f"{tenant.name} · Rapport patient confidentiel")
        canvas.drawRightString(283 * mm, 6.5 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm,
        topMargin=13 * mm, bottomMargin=16 * mm,
        title=f"Rapport patient – {data['patient_name']}",
        author=tenant.name,
    )
    story = [
        Paragraph(tenant.name, title),
        Paragraph("RAPPORT MÉDICAL ET FINANCIER DU PATIENT", ParagraphStyle("Sub", parent=styles["Normal"], textColor=colors.HexColor("#64748b"), fontSize=9)),
        Spacer(1, 4 * mm),
    ]
    identity = Table([
        [p("PATIENT", center), p("DATE DE NAISSANCE", center), p("CONTACT", center), p("GÉNÉRÉ LE", center)],
        [p(data["patient_name"]), p(patient.birth_date.strftime("%d/%m/%Y")), p(f"{patient.user.phone or '—'} · {patient.user.email or '—'}"), p(data["generated_at"].strftime("%d/%m/%Y %H:%M UTC"))],
    ], colWidths=[70 * mm, 48 * mm, 92 * mm, 55 * mm])
    identity.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b3441")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#f1f7f5")), ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#cfe2dd")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    summary = Table([
        ["TOTAL HONORAIRES", "TOTAL ENCAISSÉ", "RAP", "ACTIVITÉ MÉDICALE"],
        [_money(data["total_fees"]), _money(data["total_paid"]), _money(data["rap"]), f"{len(data['appointments'])} RDV · {len(data['visits'])} visite(s) · {len(data['prescriptions'])} ordonnance(s)"],
    ], colWidths=[66.25 * mm] * 4, rowHeights=[8 * mm, 13 * mm])
    summary.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#eaf5f2")),
        ("GRID", (0, 0), (-1, -1), .4, colors.HexColor("#cfe2dd")),
    ]))
    story += [identity, Spacer(1, 5 * mm), summary, Paragraph("Honoraires", section)]

    fee_table = [[p("Date", center), p("Libellé"), p("Statut", center), p("Montant", right), p("Payé", right), p("RAP", right)]]
    for row in data["fees"]:
        fee = row["item"]
        fee_table.append([p(fee.created_at.strftime("%d/%m/%Y"), center), p(fee.title), p(fee.status, center), p(_money(row["total"]), right), p(_money(row["paid"]), right), p(_money(row["remaining"]), right)])
    if len(fee_table) == 1:
        fee_table.append([p("—"), p("Aucun honoraire"), p("—"), p("—"), p("—"), p("—")])
    table = Table(fee_table, repeatRows=1, colWidths=[27 * mm, 94 * mm, 34 * mm, 36 * mm, 36 * mm, 36 * mm])
    table.setStyle(_table_style())
    story += [table, Paragraph("Historique des paiements", section)]
    payment_table = [[p("Date", center), p("Mode"), p("Référence"), p("Statut", center), p("Chèque", center), p("Montant", right), p("Pris en compte", center)]]
    for item in data["payments"]:
        payment_table.append([p(item.paid_at.strftime("%d/%m/%Y"), center), p(PAYMENT_LABELS.get(item.method, item.method)), p(item.reference or item.cheque_number), p(item.status, center), p(item.cheque_status, center), p(_money(item.amount), right), p("Oui" if _valid_payment(item) else "Non", center)])
    if len(payment_table) == 1:
        payment_table.append([p("—"), p("Aucun paiement"), p("—"), p("—"), p("—"), p("—"), p("—")])
    table = Table(payment_table, repeatRows=1, colWidths=[27 * mm, 36 * mm, 47 * mm, 35 * mm, 35 * mm, 38 * mm, 45 * mm])
    table.setStyle(_table_style())
    story += [table, PageBreak(), Paragraph("Ordonnances et médicaments", section)]

    prescription_rows = [[p("Date", center), p("Ordonnance"), p("Médecin"), p("Médicament"), p("Posologie / fréquence / durée"), p("Instructions")]]
    for prescription in data["prescriptions"]:
        items = prescription.items or []
        if not items:
            prescription_rows.append([p((prescription.consultation_date or prescription.created_at.date()).strftime("%d/%m/%Y"), center), p(prescription.prescription_number), p(data["doctors"].get(prescription.doctor_id, "Médecin")), p("—"), p(prescription.content), p("—")])
        for medicine in items:
            prescription_rows.append([
                p((prescription.consultation_date or prescription.created_at.date()).strftime("%d/%m/%Y"), center), p(prescription.prescription_number),
                p(data["doctors"].get(prescription.doctor_id, "Médecin")), p(medicine.get("medication_name") or medicine.get("name") or "Médicament"),
                p(" · ".join(str(value) for value in [medicine.get("dosage"), medicine.get("frequency"), medicine.get("duration")] if value)), p(medicine.get("instructions")),
            ])
    if len(prescription_rows) == 1:
        prescription_rows.append([p("—"), p("Aucune ordonnance"), p("—"), p("—"), p("—"), p("—")])
    table = Table(prescription_rows, repeatRows=1, colWidths=[24 * mm, 39 * mm, 43 * mm, 48 * mm, 69 * mm, 42 * mm])
    table.setStyle(_table_style())
    story += [table, Paragraph("Rendez-vous et visites", section)]
    care_rows = [[p("Date", center), p("Type", center), p("Médecin"), p("Motif"), p("Statut", center), p("Observations")]]
    for item in data["appointments"]:
        care_rows.append([p(item.date.strftime("%d/%m/%Y"), center), p("Rendez-vous", center), p(data["doctors"].get(item.doctor_id, "Médecin")), p(item.reason), p(item.status, center), p(item.notes)])
    for item in data["visits"]:
        care_rows.append([p(item.visit_date.strftime("%d/%m/%Y"), center), p("Visite", center), p(data["doctors"].get(item.doctor_id, "Médecin")), p(item.reason), p(item.status, center), p(item.notes)])
    care_rows[1:] = sorted(care_rows[1:], key=lambda row: row[0].text, reverse=True)
    if len(care_rows) == 1:
        care_rows.append([p("—"), p("Aucune activité"), p("—"), p("—"), p("—"), p("—")])
    table = Table(care_rows, repeatRows=1, colWidths=[25 * mm, 29 * mm, 45 * mm, 52 * mm, 34 * mm, 80 * mm])
    table.setStyle(_table_style())
    story += [table, Paragraph("Documents et dossiers CNSS/CNOPS", section)]
    admin_rows = [[p("Date", center), p("Catégorie", center), p("Libellé"), p("Statut", center), p("Informations")]]
    for item in data["documents"]:
        admin_rows.append([p(item.created_at.strftime("%d/%m/%Y"), center), p("Document", center), p(item.title), p(item.status, center), p(item.document_type)])
    for item in data["cnss_records"]:
        admin_rows.append([p(item.created_at.strftime("%d/%m/%Y"), center), p("CNSS/CNOPS", center), p(item.label), p(item.status, center), p(f"Demandé : {_money(item.requested_amount)} · Remboursé : {_money(item.reimbursed_amount or 0)}")])
    if len(admin_rows) == 1:
        admin_rows.append([p("—"), p("—"), p("Aucun document administratif"), p("—"), p("—")])
    table = Table(admin_rows, repeatRows=1, colWidths=[27 * mm, 38 * mm, 75 * mm, 40 * mm, 85 * mm])
    table.setStyle(_table_style())
    story += [table, Spacer(1, 5 * mm), Paragraph("Document confidentiel réservé aux professionnels autorisés du cabinet.", ParagraphStyle("Notice", parent=cell, textColor=colors.HexColor("#64748b"), alignment=TA_CENTER))]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b3441")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#dbe8e5")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ])


def render_patient_excel(data: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Synthèse"
    brand, dark, light = "0F766E", "0B3441", "EAF5F2"
    thin = Side(style="thin", color="D9E6E3")

    summary.append([data["tenant"].name])
    summary.append(["Rapport médical et financier du patient"])
    summary.append(["Patient", data["patient_name"]])
    summary.append(["Date de naissance", data["patient"].birth_date])
    summary.append(["Téléphone", data["patient"].user.phone or "—"])
    summary.append(["Email", data["patient"].user.email or "—"])
    summary.append([])
    summary.append(["Indicateur", "Valeur"])
    summary.append(["Total honoraires", float(data["total_fees"])])
    summary.append(["Total encaissé", float(data["total_paid"])])
    summary.append(["RAP", float(data["rap"])])
    summary.append(["Rendez-vous", len(data["appointments"])])
    summary.append(["Visites", len(data["visits"])])
    summary.append(["Ordonnances", len(data["prescriptions"])])
    summary["A1"].font = Font(size=18, bold=True, color=brand)
    summary["A2"].font = Font(size=13, bold=True, color=dark)
    _style_header(summary[8], brand)
    for coordinate in ("B9", "B10", "B11"):
        summary[coordinate].number_format = '#,##0.00 "MAD"'
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 45

    fees = workbook.create_sheet("Honoraires et paiements")
    fees.append(["HONORAIRES"])
    fees.append(["Date", "Libellé", "Statut", "Montant (MAD)", "Payé valide (MAD)", "RAP (MAD)"])
    for row in data["fees"]:
        item = row["item"]
        fees.append([item.created_at.date(), item.title, item.status, float(row["total"]), float(row["paid"]), float(row["remaining"])])
    fees.append([])
    fees.append(["PAIEMENTS — les opérations annulées restent visibles mais sont exclues du total encaissé"])
    fees.append(["Date", "Mode", "Référence", "Statut", "Statut chèque", "Montant (MAD)", "Pris en compte"])
    for item in data["payments"]:
        fees.append([item.paid_at.date(), PAYMENT_LABELS.get(item.method, item.method), item.reference or item.cheque_number or "—", item.status, item.cheque_status or "—", float(_decimal(item.amount)), "Oui" if _valid_payment(item) else "Non"])

    prescriptions = workbook.create_sheet("Ordonnances - médicaments")
    prescriptions.append(["Date", "N° ordonnance", "Médecin", "Médicament", "Dosage", "Fréquence", "Durée", "Instructions", "Texte ordonnance"])
    for item in data["prescriptions"]:
        medicines = item.items or [{}]
        for medicine in medicines:
            prescriptions.append([(item.consultation_date or item.created_at.date()), item.prescription_number, data["doctors"].get(item.doctor_id, "Médecin"), medicine.get("medication_name") or medicine.get("name") or "—", medicine.get("dosage") or "—", medicine.get("frequency") or "—", medicine.get("duration") or "—", medicine.get("instructions") or "—", item.content])

    care = workbook.create_sheet("Rendez-vous - visites")
    care.append(["Date", "Type", "Médecin", "Motif", "Statut", "Observations"])
    for item in data["appointments"]:
        care.append([item.date, "Rendez-vous", data["doctors"].get(item.doctor_id, "Médecin"), item.reason or "—", item.status, item.notes or "—"])
    for item in data["visits"]:
        care.append([item.visit_date, "Visite", data["doctors"].get(item.doctor_id, "Médecin"), item.reason or "—", item.status, item.notes or "—"])

    administration = workbook.create_sheet("Documents - CNSS-CNOPS")
    administration.append(["Date", "Catégorie", "Libellé", "Type / référence", "Statut", "Montant demandé", "Montant remboursé"])
    for item in data["documents"]:
        administration.append([item.created_at.date(), "Document", item.title, item.document_type, item.status, None, None])
    for item in data["cnss_records"]:
        administration.append([item.created_at.date(), "CNSS/CNOPS", item.label, item.reference or "—", item.status, float(_decimal(item.requested_amount)), float(_decimal(item.reimbursed_amount))])

    for sheet in workbook.worksheets[1:]:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)
        for row_number in (1, 2) if sheet is fees else (1,):
            if row_number <= sheet.max_row:
                _style_header(sheet[row_number], dark)
        if sheet is fees:
            for row_number in range(3, 3 + len(data["fees"])):
                for cell in fees[row_number][3:6]: cell.number_format = '#,##0.00 "MAD"'
            payment_header = 5 + len(data["fees"])
            _style_header(fees[payment_header], dark)
            for row_number in range(payment_header + 1, fees.max_row + 1): fees.cell(row_number, 6).number_format = '#,##0.00 "MAD"'
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for index, column in enumerate(sheet.columns, 1):
            width = min(55, max(13, max((len(str(cell.value or "")) for cell in column), default=10) + 2))
            sheet.column_dimensions[get_column_letter(index)].width = width
    summary.sheet_view.showGridLines = False
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_header(cells, color: str) -> None:
    for cell in cells:
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
