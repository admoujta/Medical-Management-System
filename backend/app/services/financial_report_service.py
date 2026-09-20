from __future__ import annotations

import calendar
import io
from datetime import UTC, date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session

from app.models import CabinetBankAccount, Expense, Payment, Tenant


MONTHS_FR = (
    "", "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
)
PAYMENT_LABELS = {
    "cash": "Espèces",
    "card": "Carte bancaire",
    "transfer": "Virement",
    "cheque": "Chèque",
}
INVALID_PAYMENT_STATUSES = {
    "ANNULE", "REMBOURSE", "ECHEC", "FAILED", "REVERSED", "CANCELLED",
}


def _money(value: Decimal | float | int) -> str:
    return f"{Decimal(str(value)):,.2f} MAD".replace(",", " ").replace(".", ",")


def _period(year: int, month: int) -> tuple[date, date, datetime, datetime]:
    last_day = calendar.monthrange(year, month)[1]
    start_date = date(year, month, 1)
    end_date = date(year, month, last_day)
    start_dt = datetime(year, month, 1, tzinfo=UTC)
    end_dt = datetime(year + (month == 12), 1 if month == 12 else month + 1, 1, tzinfo=UTC)
    return start_date, end_date, start_dt, end_dt


def monthly_financial_data(db: Session, tenant_id: str, year: int, month: int) -> dict:
    start_date, end_date, start_dt, end_dt = _period(year, month)
    tenant = db.scalar(select(Tenant).where(Tenant.id == tenant_id))
    payments = list(db.scalars(
        select(Payment).where(
            Payment.tenant_id == tenant_id,
            Payment.status.not_in(INVALID_PAYMENT_STATUSES),
            or_(
                and_(
                    Payment.method != "cheque",
                    Payment.paid_at >= start_dt,
                    Payment.paid_at < end_dt,
                ),
                and_(
                    Payment.method == "cheque",
                    Payment.cheque_status == "ENCAISSE",
                    or_(
                        and_(
                            Payment.cheque_cashing_date.is_not(None),
                            Payment.cheque_cashing_date >= start_date,
                            Payment.cheque_cashing_date <= end_date,
                        ),
                        and_(
                            Payment.cheque_cashing_date.is_(None),
                            Payment.paid_at >= start_dt,
                            Payment.paid_at < end_dt,
                        ),
                    ),
                ),
            ),
        )
    ).all())
    expenses = list(db.scalars(
        select(Expense).where(
            Expense.tenant_id == tenant_id,
            Expense.expense_date >= start_date,
            Expense.expense_date <= end_date,
        )
    ).all())
    account_ids = {
        row.bank_account_id for row in [*payments, *expenses] if row.bank_account_id
    }
    accounts = {
        row.id: row
        for row in db.scalars(
            select(CabinetBankAccount).where(CabinetBankAccount.id.in_(account_ids))
        ).all()
    } if account_ids else {}

    transactions: list[dict] = []
    methods: dict[str, Decimal] = {}
    for payment in payments:
        amount = Decimal(str(payment.amount))
        method = PAYMENT_LABELS.get(payment.method, payment.method or "Non renseigné")
        methods[method] = methods.get(method, Decimal("0")) + amount
        patient = payment.patient
        patient_name = (
            f"{patient.user.first_name} {patient.user.last_name}".strip()
            if patient and patient.user else "Patient"
        )
        account = accounts.get(payment.bank_account_id)
        operation_date = (
            payment.cheque_cashing_date
            if payment.method == "cheque" and payment.cheque_cashing_date
            else payment.paid_at.date()
        )
        transactions.append({
            "date": operation_date,
            "nature": "ENTRÉE",
            "description": f"Paiement honoraire · {patient_name}",
            "third_party": patient_name,
            "method": method,
            "account": f"{account.label} · {account.bank_name}" if account else "—",
            "reference": payment.reference or payment.cheque_number or "—",
            "entry": amount,
            "exit": Decimal("0"),
        })

    expense_categories: dict[str, Decimal] = {}
    fixed_expenses = Decimal("0")
    variable_expenses = Decimal("0")
    for expense in expenses:
        amount = Decimal(str(expense.amount_ttc if expense.amount_ttc is not None else expense.amount))
        expense_categories[expense.category] = expense_categories.get(expense.category, Decimal("0")) + amount
        if expense.expense_type == "FIXE":
            fixed_expenses += amount
        else:
            variable_expenses += amount
        account = accounts.get(expense.bank_account_id)
        transactions.append({
            "date": expense.expense_date,
            "nature": "SORTIE",
            "description": expense.label,
            "third_party": expense.supplier or "—",
            "method": PAYMENT_LABELS.get(expense.payment_method or "", expense.payment_method or "—"),
            "account": f"{account.label} · {account.bank_name}" if account else "—",
            "reference": expense.reference or "—",
            "entry": Decimal("0"),
            "exit": amount,
        })

    transactions.sort(key=lambda row: (row["date"], 0 if row["nature"] == "ENTRÉE" else 1))
    total_entries = sum((row["entry"] for row in transactions), Decimal("0"))
    total_exits = sum((row["exit"] for row in transactions), Decimal("0"))
    return {
        "tenant": tenant,
        "year": year,
        "month": month,
        "period_label": f"{MONTHS_FR[month]} {year}",
        "start_date": start_date,
        "end_date": end_date,
        "transactions": transactions,
        "total_entries": total_entries,
        "total_exits": total_exits,
        "net": total_entries - total_exits,
        "payment_count": len(payments),
        "expense_count": len(expenses),
        "fixed_expenses": fixed_expenses,
        "variable_expenses": variable_expenses,
        "methods": methods,
        "expense_categories": expense_categories,
    }


def render_monthly_pdf(data: dict) -> bytes:
    output = io.BytesIO()
    tenant = data["tenant"]
    brand = colors.HexColor(getattr(tenant, "primary_color", None) or "#0f766e")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], textColor=brand, fontSize=20, leading=24)
    subtitle = ParagraphStyle("ReportSubtitle", parent=styles["Normal"], textColor=colors.HexColor("#64748b"), fontSize=9)
    cell = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=7.3, leading=9)
    money_cell = ParagraphStyle("Money", parent=cell, alignment=TA_RIGHT)
    center = ParagraphStyle("Center", parent=cell, alignment=TA_CENTER)

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#dbe8e5"))
        canvas.line(15 * mm, 12 * mm, 282 * mm, 12 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(15 * mm, 7 * mm, f"{tenant.name} · Rapport financier {data['period_label']}")
        canvas.drawRightString(282 * mm, 7 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=17 * mm,
        title=f"Rapport financier {data['period_label']} – {tenant.name}",
    )
    story = [
        Paragraph(tenant.name, title),
        Paragraph(
            f"RAPPORT FINANCIER MENSUEL · {data['period_label'].upper()}<br/>"
            f"Période du {data['start_date'].strftime('%d/%m/%Y')} au {data['end_date'].strftime('%d/%m/%Y')} · Généré le {datetime.now(UTC).strftime('%d/%m/%Y à %H:%M UTC')}",
            subtitle,
        ),
        Spacer(1, 6 * mm),
    ]
    summary = [
        ["ENTRÉES ENCAISSÉES", "SORTIES / DÉPENSES", "SOLDE NET", "OPÉRATIONS"],
        [
            _money(data["total_entries"]),
            _money(data["total_exits"]),
            _money(data["net"]),
            f"{data['payment_count']} entrée(s) · {data['expense_count']} sortie(s)",
        ],
    ]
    summary_table = Table(summary, colWidths=[66 * mm] * 4, rowHeights=[8 * mm, 13 * mm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("FONTSIZE", (0, 1), (-1, 1), 12),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, 1), (-1, 1), colors.HexColor("#f1f7f5")),
        ("BOX", (0, 0), (-1, -1), .5, colors.HexColor("#cfe2dd")),
        ("INNERGRID", (0, 0), (-1, -1), .35, colors.HexColor("#cfe2dd")),
    ]))
    story += [summary_table, Spacer(1, 6 * mm)]

    rows = [[
        Paragraph("<b>Date</b>", center), Paragraph("<b>Nature</b>", center),
        Paragraph("<b>Description</b>", cell), Paragraph("<b>Tiers</b>", cell),
        Paragraph("<b>Mode / compte</b>", cell), Paragraph("<b>Référence</b>", cell),
        Paragraph("<b>Entrée</b>", money_cell), Paragraph("<b>Sortie</b>", money_cell),
    ]]
    for row in data["transactions"]:
        rows.append([
            Paragraph(row["date"].strftime("%d/%m/%Y"), center),
            Paragraph(f"<b>{row['nature']}</b>", center),
            Paragraph(row["description"], cell),
            Paragraph(row["third_party"], cell),
            Paragraph(f"{row['method']}<br/><font color='#64748b'>{row['account']}</font>", cell),
            Paragraph(row["reference"], cell),
            Paragraph(_money(row["entry"]) if row["entry"] else "—", money_cell),
            Paragraph(_money(row["exit"]) if row["exit"] else "—", money_cell),
        ])
    if not data["transactions"]:
        rows.append(["", "", Paragraph("Aucune opération enregistrée pour ce mois.", center), "", "", "", "", ""])
    table = Table(
        rows, repeatRows=1,
        colWidths=[20 * mm, 21 * mm, 55 * mm, 37 * mm, 47 * mm, 31 * mm, 28 * mm, 28 * mm],
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0b3441")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#dbe8e5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [Paragraph("Détail des entrées et sorties", styles["Heading2"]), Spacer(1, 2 * mm), table]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render_monthly_excel(data: dict) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Synthèse"
    detail = workbook.create_sheet("Entrées - Sorties")
    brand = "0F766E"
    dark = "0B3441"
    light = "EAF5F2"
    thin = Side(style="thin", color="D9E6E3")

    summary.merge_cells("A1:D1")
    summary["A1"] = data["tenant"].name
    summary["A1"].font = Font(size=18, bold=True, color=brand)
    summary.merge_cells("A2:D2")
    summary["A2"] = f"Rapport financier mensuel · {data['period_label']}"
    summary["A2"].font = Font(size=13, bold=True, color=dark)
    summary.merge_cells("A3:D3")
    summary["A3"] = f"Du {data['start_date'].strftime('%d/%m/%Y')} au {data['end_date'].strftime('%d/%m/%Y')}"
    summary["A3"].font = Font(size=10, color="64748B")
    summary.append([])
    summary.append(["Indicateur", "Montant / nombre", "Indicateur", "Montant / nombre"])
    summary.append(["Entrées encaissées", float(data["total_entries"]), "Nombre d’entrées", data["payment_count"]])
    summary.append(["Sorties / dépenses", float(data["total_exits"]), "Nombre de sorties", data["expense_count"]])
    summary.append(["Solde net", float(data["net"]), "Charges fixes", float(data["fixed_expenses"])])
    summary.append(["", "", "Charges variables", float(data["variable_expenses"])])
    for cell in summary[5]:
        cell.fill = PatternFill("solid", fgColor=brand)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for row in summary.iter_rows(min_row=6, max_row=9, min_col=1, max_col=4):
        for cell in row:
            cell.border = Border(bottom=thin)
    for coordinate in ("B6", "B7", "B8", "D8", "D9"):
        summary[coordinate].number_format = '#,##0.00 "MAD"'
    summary.column_dimensions["A"].width = 27
    summary.column_dimensions["B"].width = 22
    summary.column_dimensions["C"].width = 25
    summary.column_dimensions["D"].width = 22

    headers = ["Date", "Nature", "Description", "Tiers", "Mode de paiement", "Compte bancaire", "Référence", "Entrée (MAD)", "Sortie (MAD)", "Solde cumulé (MAD)"]
    detail.append(headers)
    running = Decimal("0")
    for transaction in data["transactions"]:
        running += transaction["entry"] - transaction["exit"]
        detail.append([
            transaction["date"], transaction["nature"], transaction["description"],
            transaction["third_party"], transaction["method"], transaction["account"],
            transaction["reference"], float(transaction["entry"]), float(transaction["exit"]),
            float(running),
        ])
    for cell in detail[1]:
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in detail.iter_rows(min_row=2):
        row[0].number_format = "dd/mm/yyyy"
        for cell in row[7:10]:
            cell.number_format = '#,##0.00 "MAD"'
        if row[1].value == "ENTRÉE":
            row[1].font = Font(bold=True, color="047857")
        else:
            row[1].font = Font(bold=True, color="B91C1C")
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [13, 13, 38, 26, 20, 30, 20, 18, 18, 21]
    for index, width in enumerate(widths, 1):
        detail.column_dimensions[get_column_letter(index)].width = width
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:J{max(1, detail.max_row)}"
    detail.sheet_view.showGridLines = False
    summary.sheet_view.showGridLines = False
    detail.page_setup.orientation = "landscape"
    detail.page_setup.fitToWidth = 1

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
